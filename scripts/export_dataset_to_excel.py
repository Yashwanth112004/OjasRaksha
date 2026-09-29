import json
import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_excel():
    base_dir = os.path.dirname(os.path.abspath(__file__))
    json_path = os.path.join(base_dir, 'synthetic_dataset_3000.json')
    excel_path = os.path.join(base_dir, 'synthetic_dataset_3000.xlsx')
    csv_path = os.path.join(base_dir, 'synthetic_dataset_3000.csv')

    print(f"Reading JSON dataset from: {json_path}")
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    # 1. FLATTEN CLINICAL TRANSACTIONS (3,000 ROWS)
    flattened_rows = []
    for item in data:
        meds_str = "; ".join([f"{m.get('name')} ({m.get('dose')})" for m in item['clinicalDetails'].get('prescribedMedications', [])])
        labs_str = "; ".join(item['clinicalDetails'].get('diagnosticLabsRequested', []))
        
        row = {
            "Row_ID": item.get("rowId"),
            "Record_ID": item.get("recordId"),
            "Encounter_Timestamp": item.get("encounterTimestamp"),
            
            # Hospital & Applied Rules
            "Hospital_ID": item['hospital'].get("id"),
            "Hospital_Name": item['hospital'].get("name"),
            "Hospital_Type": item['hospital'].get("type"),
            "Hospital_Wallet": item['hospital'].get("walletAddress"),
            "Hospital_Consent_TTL_Days": item['hospital']['appliedRules'].get("consentDurationDays"),
            "Hospital_DPDP_Tier": item['hospital']['appliedRules'].get("dpdpTier"),
            "Break_Glass_Permitted": "YES" if item['hospital']['appliedRules'].get("breakGlassPermitted") else "NO",
            "MPC_Threshold_Schema": item['hospital']['appliedRules'].get("mpcThreshold"),
            "Data_Retention_Years": item['hospital']['appliedRules'].get("dataRetentionYears"),
            "Insurance_PreAuth_Required": "YES" if item['hospital']['appliedRules'].get("requireInsurancePreAuth") else "NO",

            # Patient Demographics
            "Patient_ID": item['patient'].get("patientId"),
            "Patient_Short_ID": item['patient'].get("shortId"),
            "Patient_Wallet": item['patient'].get("walletAddress"),
            "Patient_Age": item['patient'].get("age"),
            "Patient_Gender": item['patient'].get("gender"),
            "Patient_Masked_Phone": item['patient'].get("maskedPhone"),

            # Attending Doctor
            "Doctor_Name": item['attendingDoctor'].get("name"),
            "Doctor_Wallet": item['attendingDoctor'].get("walletAddress"),
            "Doctor_NMC_RegNo": item['attendingDoctor'].get("registrationNo"),
            "Doctor_Specialty": item['attendingDoctor'].get("specialtyRole"),

            # Clinical & ICD-10 Classification
            "Record_Type": item['clinicalDetails'].get("recordType"),
            "ICD10_Code": item['clinicalClassification'].get("icd10Code"),
            "Diagnosis_Title": item['clinicalClassification'].get("diagnosisTitle"),
            "ICD10_Chapter": item['clinicalClassification'].get("chapter"),
            "Clinical_Category": item['clinicalClassification'].get("clinicalCategory"),
            "Severity_Level": item['clinicalClassification'].get("severityLevel"),
            "Billable_Status": "Billable" if item['clinicalClassification'].get("billableStatus") else "Non-Billable",
            "Prescribed_Medications": meds_str,
            "Diagnostic_Labs_Requested": labs_str,
            "Estimated_Cost_INR": item['clinicalDetails'].get("estimatedCostINR"),

            # Cryptographic Envelope (AES-256-GCM)
            "Encryption_Cipher": item['cryptographicEnvelope'].get("algorithm"),
            "Vault_Transit_Key_Handle": item['cryptographicEnvelope'].get("transitKeyHandle"),
            "Initialization_Vector_Hex_96bit": item['cryptographicEnvelope'].get("initializationVectorHex"),
            "Galois_Auth_Tag_Hex_128bit": item['cryptographicEnvelope'].get("authTagHex"),
            "Raw_Payload_Bytes": item['cryptographicEnvelope'].get("rawPayloadBytes"),
            "Ciphertext_Length_Bytes": item['cryptographicEnvelope'].get("ciphertextLengthBytes"),
            "Fixed_Crypto_Overhead_Bytes": item['cryptographicEnvelope'].get("fixedOverheadBytes"),
            "ZK_Commitment_Proof": item['cryptographicEnvelope'].get("zkCommitmentProof"),

            # Storage & Hedera Consensus Audit
            "Storage_Network": item['decentralizedStorage'].get("storageNetwork"),
            "IPFS_CID": item['decentralizedStorage'].get("ipfsCID"),
            "HCS_Topic_ID": item['hederaConsensusAudit'].get("hcsTopicId"),
            "HCS_Sequence_Number": item['hederaConsensusAudit'].get("sequenceNumber"),
            "HCS_Consensus_Timestamp": item['hederaConsensusAudit'].get("consensusTimestamp"),
            "HCS_Running_Hash": item['hederaConsensusAudit'].get("runningHash"),
            "Finality_Latency_Sec": float(item['hederaConsensusAudit'].get("finalityLatencySeconds")),
            "Gas_Fee_HBAR": float(item['hederaConsensusAudit'].get("gasFeeHbar")),

            # DPDP Act 2023 Governance
            "DPDP_Consent_Status": item['dpdpGovernance'].get("consentStatus"),
            "Purpose_Limitation": item['dpdpGovernance'].get("purposeLimitation"),
            "Retention_Expiry_Date": item['dpdpGovernance'].get("retentionExpiryDate"),
            "Right_To_Erasure_Supported": "YES" if item['dpdpGovernance'].get("rightToErasureSupported") else "NO"
        }
        flattened_rows.append(row)

    df_main = pd.DataFrame(flattened_rows)

    # Export to CSV as well
    df_main.to_csv(csv_path, index=False, encoding='utf-8')
    print(f"Exported CSV dataset to: {csv_path}")

    # 2. HOSPITAL RULES MATRIX SHEET
    hospital_rules_list = [
        {
            "Hospital_ID": "HOSP-001",
            "Hospital_Name": "Apollo Multispecialty Hospital",
            "Type": "Tertiary Multi-Specialty",
            "Wallet_Address": "0x155Af6ECaFb48861dA7d16Fb8Af2f6ce9d6DD779",
            "Specialties": "Cardiology, Infectious Disease, Gastroenterology, General Medicine",
            "Consent_Duration_Days": 365,
            "DPDP_Tier": "Standard Sovereign Tier-1",
            "Break_Glass_Permitted": "YES (Critical & Emergency Only)",
            "MPC_Threshold": "2-of-3 Web3Auth Enclave",
            "Encryption_Cipher": "AES-256-GCM (128-bit AuthTag)",
            "Data_Retention_Years": 7,
            "Merkle_Batch_Size": 50,
            "Telecom_Masking": "Partial (Last 4 digits visible)",
            "Insurance_PreAuth_Required": "YES"
        },
        {
            "Hospital_ID": "HOSP-002",
            "Hospital_Name": "Max Super Speciality Cancer Institute",
            "Type": "Comprehensive Oncology Care",
            "Wallet_Address": "0x28B7eA94B34C8F52D19E50Eb14769Da20F5D54b4",
            "Specialties": "Oncology, Hematology & Immunology, Biochemical Genetics",
            "Consent_Duration_Days": 180,
            "DPDP_Tier": "Sensitive Biomarker / Genomic Tier",
            "Break_Glass_Permitted": "NO (Ethics Board Signoff Required)",
            "MPC_Threshold": "3-of-3 MPC Full Strict",
            "Encryption_Cipher": "AES-256-GCM (128-bit AuthTag)",
            "Data_Retention_Years": 15,
            "Merkle_Batch_Size": 25,
            "Telecom_Masking": "Full Cryptographic Anonymization",
            "Insurance_PreAuth_Required": "YES"
        },
        {
            "Hospital_ID": "HOSP-003",
            "Hospital_Name": "Fortis Memorial Neuro & Trauma Center",
            "Type": "Level-1 Emergency & Neurological Institute",
            "Wallet_Address": "0x84f93618C1D541bB1Dcf1373A8eE500A7D439634",
            "Specialties": "Neurology, Trauma & Emergency, Cardiology",
            "Consent_Duration_Days": 90,
            "DPDP_Tier": "Emergency Rapid Fast-Track",
            "Break_Glass_Permitted": "YES (60-min Ephemeral Auto-Expiry)",
            "MPC_Threshold": "2-of-3 Ephemeral Enclave",
            "Encryption_Cipher": "AES-256-GCM (128-bit AuthTag)",
            "Data_Retention_Years": 10,
            "Merkle_Batch_Size": 50,
            "Telecom_Masking": "Emergency Unmasking Authorized",
            "Insurance_PreAuth_Required": "NO (Emergency Fast-Track)"
        },
        {
            "Hospital_ID": "HOSP-004",
            "Hospital_Name": "Care & Cure Pediatric & Maternity Hospital",
            "Type": "Mother & Child Specialized Center",
            "Wallet_Address": "0x937F51E09477BfD31E56c5e533038a8eEAc096b7",
            "Specialties": "Pediatrics, Obstetrics & Gynecology, Genetics, Neonatology",
            "Consent_Duration_Days": 365,
            "DPDP_Tier": "DPDP Section 9 Child Data Protection",
            "Break_Glass_Permitted": "YES (Neonatal & Obstetric Emergency)",
            "MPC_Threshold": "2-of-3 Guardian & Hospital Enclave",
            "Encryption_Cipher": "AES-256-GCM (128-bit AuthTag)",
            "Data_Retention_Years": 21,
            "Merkle_Batch_Size": 50,
            "Telecom_Masking": "Guardian Telecom Linked",
            "Insurance_PreAuth_Required": "NO"
        },
        {
            "Hospital_ID": "HOSP-005",
            "Hospital_Name": "Apex Diagnostic & Pathology Reference Labs",
            "Type": "Centralized Diagnostic Reference Hub",
            "Wallet_Address": "0x04Fee3FD1B338d12FFD6dBD8d66dE1e8e0BB99cB",
            "Specialties": "Endocrinology, Diagnostics, Infectious Disease, Hematology",
            "Consent_Duration_Days": 30,
            "DPDP_Tier": "Diagnostic Telemetry Tier",
            "Break_Glass_Permitted": "NO (Telemetry Push Only)",
            "MPC_Threshold": "2-of-3 Node Share Enclave",
            "EncryptionCipher": "AES-256-GCM (128-bit AuthTag)",
            "Data_Retention_Years": 5,
            "Merkle_Batch_Size": 100,
            "Telecom_Masking": "Full Masking",
            "Insurance_PreAuth_Required": "YES"
        }
    ]
    df_rules = pd.DataFrame(hospital_rules_list)

    # 3. STATISTICAL SUMMARY SHEET
    hosp_summary = df_main.groupby('Hospital_Name').size().reset_index(name='Record_Count')
    hosp_summary['Percentage'] = (hosp_summary['Record_Count'] / len(df_main) * 100).round(2).astype(str) + '%'

    cat_summary = df_main.groupby('Clinical_Category').size().reset_index(name='Record_Count')
    cat_summary['Percentage'] = (cat_summary['Record_Count'] / len(df_main) * 100).round(2).astype(str) + '%'

    consent_summary = df_main.groupby('DPDP_Consent_Status').size().reset_index(name='Record_Count')
    consent_summary['Percentage'] = (consent_summary['Record_Count'] / len(df_main) * 100).round(2).astype(str) + '%'

    # WRITE MULTI-SHEET EXCEL WORKBOOK
    with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
        df_main.to_excel(writer, sheet_name='Clinical_Dataset_3000', index=False)
        df_rules.to_excel(writer, sheet_name='Hospital_Governance_Rules', index=False)
        hosp_summary.to_excel(writer, sheet_name='Hospital_Distribution', index=False)
        cat_summary.to_excel(writer, sheet_name='Clinical_Category_Distribution', index=False)
        consent_summary.to_excel(writer, sheet_name='Consent_Status_Distribution', index=False)

    # 4. EXCEL STYLING (Headers, Borders, Auto-fit Columns)
    wb = openpyxl.load_workbook(excel_path)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        ws.views.sheetView[0].showGridLines = True
        
        # Style Header Row
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Style Data Cells & Auto-fit width
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row > 1:
                    cell.border = thin_border
                    cell.alignment = Alignment(vertical="center")
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 40)

        ws.row_dimensions[1].height = 28

    wb.save(excel_path)
    print(f"[SUCCESS] Multi-sheet styled Excel workbook generated at: {excel_path}")

if __name__ == "__main__":
    generate_excel()
