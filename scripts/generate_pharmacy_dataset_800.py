import os
import json
import random
import hashlib
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def generate_pharmacy_dataset(total_rows=800):
    print("================================================================================")
    print(f"GENERATING OJASRAKSHA PHARMACY PRESCRIPTION QUEUE DATASET ({total_rows} ROWS)")
    print("================================================================================\n")

    base_dir = os.path.dirname(os.path.abspath(__file__))
    json_path = os.path.join(base_dir, 'pharmacy_dataset_800.json')
    csv_path = os.path.join(base_dir, 'pharmacy_dataset_800.csv')
    excel_path = os.path.join(base_dir, 'pharmacy_dataset_800.xlsx')

    # Root workspace copy paths
    root_dir = os.path.abspath(os.path.join(base_dir, '..', '..'))
    root_excel_path = os.path.join(root_dir, 'pharmacy_dataset_800.xlsx')
    root_csv_path = os.path.join(root_dir, 'pharmacy_dataset_800.csv')

    # 1. PHARMACIES & DISPENSARY HUBS
    PHARMACIES = [
        {"id": "PHARM-001", "name": "Apollo Pharmacy Global", "wallet": "0x04Fee3FD1B338d12FFD6dBD8d66dE1e8e0BB99cB", "license": "DL-AP-2024-8901", "leadPharmacist": "K. V. Raman, R.Ph."},
        {"id": "PHARM-002", "name": "MedPlus Health Hub", "wallet": "0x70997970C51812dc3A010C7d01b50e0d17dc79C8", "license": "DL-MP-2023-4122", "leadPharmacist": "Priya Sundaram, Pharm.D."},
        {"id": "PHARM-003", "name": "Fortis Care Dispensary", "wallet": "0x84f93618C1D541bB1Dcf1373A8eE500A7D439634", "license": "DL-FC-2025-1094", "leadPharmacist": "Amit Singhania, R.Ph."},
        {"id": "PHARM-004", "name": "Max Wellness Pharmacy", "wallet": "0x28B7eA94B34C8F52D19E50Eb14769Da20F5D54b4", "license": "DL-MW-2024-7731", "leadPharmacist": "Sneha Mukherjee, R.Ph."},
        {"id": "PHARM-005", "name": "Apex Central Clinical Dispensary", "wallet": "0x155Af6ECaFb48861dA7d16Fb8Af2f6ce9d6DD779", "license": "DL-AX-2026-3021", "leadPharmacist": "Rajesh Nambiar, Pharm.D."}
    ]

    # 2. PATIENT POOL (Diverse Indian & Global Names)
    FIRST_NAMES = [
        "Aarav", "Aditi", "Ananya", "Arjun", "Deepak", "Divya", "Gautam", "Ishaan", "Kavya", "Manish",
        "Meera", "Neha", "Nikhil", "Pooja", "Pranav", "Priyanka", "Rahul", "Rakesh", "Riya", "Rohan",
        "Sakshi", "Sameer", "Sanjay", "Shreya", "Siddharth", "Sneha", "Tanvi", "Varun", "Vikas", "Yashwanth",
        "Amit", "Anand", "Bhavna", "Chetan", "Geeta", "Harish", "Kiran", "Lata", "Madhav", "Nandini"
    ]
    LAST_NAMES = [
        "Sharma", "Verma", "Patel", "Reddy", "Gupta", "Nair", "Iyer", "Rao", "Deshmukh", "Choudhury",
        "Mehta", "Bhatia", "Joshi", "Kapoor", "Banerjee", "Chatterjee", "Mishra", "Pandey", "Saxena", "Agarwal",
        "Kulkarni", "Shetty", "Pillai", "Menon", "Gowda", "Hegde", "Ghosh", "Mukherjee", "Swaminathan", "Malhotra"
    ]

    # 3. DOCTORS & HOSPITALS
    DOCTORS = [
        {"name": "Dr. Sarah Jenkins", "wallet": "0x04Fee3FD1B338d12FFD6dBD8d66dE1e8e0BB99cB", "hospital": "Apollo Multispecialty Hospital", "regNo": "NMC-48192", "specialty": "Cardiology"},
        {"name": "Dr. Rajesh Sharma", "wallet": "0x3c44CdDdB6a900fa2b585dd299e03d12FA4293BC", "hospital": "Apollo Multispecialty Hospital", "regNo": "NMC-78491", "specialty": "Gastroenterology"},
        {"name": "Dr. Vikram Seth", "wallet": "0x15d34AAf54267DB7D7c367839AAf71A00a2C6A65", "hospital": "Max Super Speciality Cancer Institute", "regNo": "NMC-99014", "specialty": "Oncology"},
        {"name": "Dr. Ananya Roy", "wallet": "0x9965507D1a55bcC2695C58ba16FB37d819B0A4dc", "hospital": "Max Super Speciality Cancer Institute", "regNo": "NMC-51203", "specialty": "Hematology"},
        {"name": "Dr. Rohan Patel", "wallet": "0x976EA74026E726554dB657fA54763abd0C3a0aa9", "hospital": "Fortis Memorial Neuro & Trauma Center", "regNo": "NMC-11029", "specialty": "Neurology"},
        {"name": "Dr. Preeti Deshmukh", "wallet": "0x23618e81E3f5cdF7f54C3d65f7FBc0aBf5B21E8f", "hospital": "Care & Cure Pediatric & Maternity Hospital", "regNo": "NMC-82194", "specialty": "Obstetrics & Pediatrics"},
        {"name": "Dr. Sunita Narang", "wallet": "0xBcd4042DE499D14e55001CcbB24a551F3b954096", "hospital": "Apex Diagnostic Reference Labs", "regNo": "NMC-90812", "specialty": "Endocrinology & Internal Medicine"}
    ]

    # 4. CLINICAL DRUG CATALOG WITH ICD-10 MAPPINGS
    PRESCRIPTION_CATALOG = [
        {
            "icd10": "I10", "title": "Essential (primary) hypertension", "category": "Cardiology",
            "meds": [
                {"name": "Amlodipine Besylate", "dosage": "5 mg", "freq": "Once Daily (Morning)", "days": 30, "qty": 30, "unitPrice": 6.50},
                {"name": "Telmisartan", "dosage": "40 mg", "freq": "Once Daily (Morning)", "days": 30, "qty": 30, "unitPrice": 9.20}
            ]
        },
        {
            "icd10": "I20.9", "title": "Angina pectoris, unspecified", "category": "Cardiology",
            "meds": [
                {"name": "Nitroglycerin Sublingual", "dosage": "0.4 mg", "freq": "PRN for chest pain", "days": 30, "qty": 25, "unitPrice": 14.00},
                {"name": "Aspirin Gastro-resistant", "dosage": "75 mg", "freq": "Once Daily post meal", "days": 30, "qty": 30, "unitPrice": 3.80},
                {"name": "Atorvastatin Calcium", "dosage": "40 mg", "freq": "Once Daily at Bedtime", "days": 30, "qty": 30, "unitPrice": 16.50}
            ]
        },
        {
            "icd10": "E11.9", "title": "Type 2 diabetes mellitus without complications", "category": "Endocrinology",
            "meds": [
                {"name": "Metformin Hydrochloride", "dosage": "500 mg", "freq": "Twice Daily with meals", "days": 30, "qty": 60, "unitPrice": 4.20},
                {"name": "Dapagliflozin Propanediol", "dosage": "10 mg", "freq": "Once Daily morning", "days": 30, "qty": 30, "unitPrice": 22.00},
                {"name": "Glimepiride", "dosage": "2 mg", "freq": "Once Daily pre-breakfast", "days": 30, "qty": 30, "unitPrice": 8.10}
            ]
        },
        {
            "icd10": "E11.65", "title": "Type 2 diabetes mellitus with hyperglycemia", "category": "Endocrinology",
            "meds": [
                {"name": "Insulin Glargine Solostar", "dosage": "100 IU/mL", "freq": "14 Units Once Daily at Bedtime", "days": 30, "qty": 2, "unitPrice": 680.00},
                {"name": "Metformin XR", "dosage": "1000 mg", "freq": "Once Daily post dinner", "days": 30, "qty": 30, "unitPrice": 7.50}
            ]
        },
        {
            "icd10": "E03.9", "title": "Hypothyroidism, unspecified", "category": "Endocrinology",
            "meds": [
                {"name": "Levothyroxine Sodium", "dosage": "50 mcg", "freq": "Once Daily on empty stomach", "days": 60, "qty": 60, "unitPrice": 2.90}
            ]
        },
        {
            "icd10": "G43.909", "title": "Migraine, unspecified, not intractable", "category": "Neurology",
            "meds": [
                {"name": "Sumatriptan Succinate", "dosage": "50 mg", "freq": "1 tablet at onset of migraine SOS", "days": 30, "qty": 10, "unitPrice": 45.00},
                {"name": "Naproxen Sodium", "dosage": "500 mg", "freq": "Twice daily as needed with food", "days": 10, "qty": 20, "unitPrice": 9.50},
                {"name": "Propranolol HCl", "dosage": "40 mg", "freq": "Twice daily for prophylaxis", "days": 30, "qty": 60, "unitPrice": 5.20}
            ]
        },
        {
            "icd10": "G40.909", "title": "Epilepsy, unspecified", "category": "Neurology",
            "meds": [
                {"name": "Levetiracetam", "dosage": "500 mg", "freq": "Twice Daily morning & evening", "days": 30, "qty": 60, "unitPrice": 18.00},
                {"name": "Clobazam", "dosage": "10 mg", "freq": "Once Daily at Bedtime", "days": 30, "qty": 30, "unitPrice": 12.50}
            ]
        },
        {
            "icd10": "C50.919", "title": "Malignant neoplasm of female breast", "category": "Oncology",
            "meds": [
                {"name": "Tamoxifen Citrate", "dosage": "20 mg", "freq": "Once Daily at fixed hour", "days": 30, "qty": 30, "unitPrice": 38.00},
                {"name": "Ondansetron HCl", "dosage": "8 mg", "freq": "Twice Daily 30 min before food", "days": 15, "qty": 30, "unitPrice": 11.20},
                {"name": "Calcium Carbonate + Vit D3", "dosage": "500mg/250IU", "freq": "Once Daily post meal", "days": 30, "qty": 30, "unitPrice": 6.00}
            ]
        },
        {
            "icd10": "C61", "title": "Malignant neoplasm of prostate", "category": "Oncology",
            "meds": [
                {"name": "Bicalutamide", "dosage": "50 mg", "freq": "Once Daily morning", "days": 30, "qty": 30, "unitPrice": 95.00},
                {"name": "Tamsulosin HCl", "dosage": "0.4 mg", "freq": "Once Daily 30 min after same meal", "days": 30, "qty": 30, "unitPrice": 14.50}
            ]
        },
        {
            "icd10": "K21.9", "title": "Gastro-esophageal reflux disease", "category": "Gastroenterology",
            "meds": [
                {"name": "Pantoprazole Gastro-resistant", "dosage": "40 mg", "freq": "Once Daily 30 min pre-breakfast", "days": 30, "qty": 30, "unitPrice": 8.00},
                {"name": "Domperidone Sustained Release", "dosage": "30 mg", "freq": "Once Daily morning", "days": 30, "qty": 30, "unitPrice": 11.50},
                {"name": "Sucralfate Oral Suspension", "dosage": "1000 mg/10mL", "freq": "10 mL thrice daily before meals", "days": 15, "qty": 2, "unitPrice": 165.00}
            ]
        },
        {
            "icd10": "J06.9", "title": "Acute upper respiratory infection", "category": "Pediatrics / General",
            "meds": [
                {"name": "Amoxicillin + Potassium Clavulanate", "dosage": "625 mg", "freq": "Twice Daily post food", "days": 7, "qty": 14, "unitPrice": 22.50},
                {"name": "Paracetamol", "dosage": "650 mg", "freq": "Thrice daily post meals SOS", "days": 5, "qty": 15, "unitPrice": 2.50},
                {"name": "Levocetirizine Dihydrochloride", "dosage": "5 mg", "freq": "Once Daily at Bedtime", "days": 7, "qty": 7, "unitPrice": 6.00}
            ]
        },
        {
            "icd10": "A90", "title": "Dengue fever [classical dengue]", "category": "Infectious Disease",
            "meds": [
                {"name": "Paracetamol Infusion / Tab", "dosage": "650 mg", "freq": "Every 6 hours SOS fever > 100°F", "days": 7, "qty": 20, "unitPrice": 3.00},
                {"name": "Carica Papaya Leaf Extract", "dosage": "1100 mg", "freq": "Thrice Daily post food", "days": 7, "qty": 21, "unitPrice": 32.00},
                {"name": "Oral Rehydration Salts (WHO formula)", "dosage": "21.8 g sachet", "freq": "1 sachet in 1L water daily", "days": 7, "qty": 7, "unitPrice": 22.00}
            ]
        }
    ]

    raw_records = []
    base_epoch = 1789300000000 # Sept 2026

    for i in range(1, total_rows + 1):
        # 1. Patient Info
        patient_name = f"{random.choice(FIRST_NAMES)} {random.choice(LAST_NAMES)}"
        patient_short_id = str(849200 + (i % 450))
        patient_wallet = "0x" + hashlib.sha256(f"PATIENT_PHARM_{patient_short_id}".encode()).hexdigest()[:40]
        
        # 2. Doctor & Hospital
        doctor = random.choice(DOCTORS)
        pharmacy = random.choice(PHARMACIES)

        # 3. Clinical Prescription Details
        diag_item = random.choice(PRESCRIPTION_CATALOG)
        medications = diag_item["meds"]

        # Calculate exact Bill Amount INR
        total_bill_inr = sum([m["qty"] * m["unitPrice"] for m in medications])
        total_bill_inr = round(total_bill_inr + random.uniform(20.0, 50.0), 2) # Adding standard dispensary dispensing fee

        # Meds summary text
        med_summary = " | ".join([f"{m['name']} ({m['dosage']}, {m['freq']}, Qty: {m['qty']})" for m in medications])

        # 4. Action Status & Authorization Logic (Matching Frontend Portal UI)
        # States: 
        #   - "🔓 Decrypt RX" (isAuthorized = True, Consent Granted)
        #   - "⏳ Waiting for Approval" (Access Requested, pending on-chain)
        #   - "Request Access" (Not yet requested)
        #   - "Dispensed" (Already fulfilled on-chain)
        status_rand = random.random()
        if status_rand < 0.45:
            dispensation_status = "Pending - Ready to Decrypt"
            ui_action_button = "🔓 Decrypt RX"
            is_authorized = "YES (Patient Consent Active)"
            has_pending_req = "NO"
            is_dispensed_on_chain = "NO"
            dispensed_timestamp = "N/A"
        elif status_rand < 0.70:
            dispensation_status = "Dispensed & Fulfilled"
            ui_action_button = "Dispensed (Fulfilled)"
            is_authorized = "YES"
            has_pending_req = "NO"
            is_dispensed_on_chain = "YES"
            dispensed_timestamp = f"{i*120}s post-verification"
        elif status_rand < 0.88:
            dispensation_status = "Access Pending Approval"
            ui_action_button = "⏳ Waiting for Approval"
            is_authorized = "NO"
            has_pending_req = "YES"
            is_dispensed_on_chain = "NO"
            dispensed_timestamp = "N/A"
        else:
            dispensation_status = "Unauthorized / Unlinked"
            ui_action_button = "Request Access"
            is_authorized = "NO"
            has_pending_req = "NO"
            is_dispensed_on_chain = "NO"
            dispensed_timestamp = "N/A"

        # 5. IPFS & Hedera Blockchain Receipt
        cid_hash = hashlib.sha256(f"PRESCRIPTION_IPFS_{i}_{patient_wallet}_{diag_item['icd10']}".encode()).hexdigest()
        ipfs_cid = f"Qm{cid_hash[:44]}"
        
        hcs_seq = 2000 + i
        time_offset = (i * 2400) + random.randint(100, 900)
        consensus_timestamp = pd.to_datetime(base_epoch + time_offset, unit='ms').strftime('%Y-%m-%d %H:%M:%S UTC')
        running_hash = "0x" + hashlib.sha256(f"HCS_RUNNING_{hcs_seq}_{ipfs_cid}".encode()).hexdigest()

        # 6. AES-256-GCM Cryptographic Parameters
        iv_hex = hashlib.sha256(f"IV_{i}_{patient_short_id}".encode()).hexdigest()[:24] # 96-bit
        auth_tag_hex = hashlib.sha256(f"AUTHTAG_{i}_{patient_short_id}".encode()).hexdigest()[:32] # 128-bit

        record = {
            # Core UI Columns from Screenshot
            "Record_ID": f"RX-{str(i).zfill(4)}",
            "Patient_Identity": patient_name,
            "Short_ID": patient_short_id,
            "Wallet_Address": patient_wallet,
            "Action_Button_UI": ui_action_button,
            
            # Queue & Fulfillment Status
            "Dispensation_Status": dispensation_status,
            "Is_Authorized_By_Consent": is_authorized,
            "Pending_Access_Request": has_pending_req,
            "Is_Dispensed_On_Chain": is_dispensed_on_chain,
            "Total_Bill_Amount_INR": total_bill_inr,
            
            # Prescribing Clinician & Hospital
            "Prescribing_Doctor": doctor["name"],
            "Doctor_Wallet": doctor["wallet"],
            "Doctor_NMC_Reg": doctor["regNo"],
            "Originating_Hospital": doctor["hospital"],

            # Assigned Pharmacy & Pharmacist
            "Assigned_Pharmacy": pharmacy["name"],
            "Pharmacy_Wallet": pharmacy["wallet"],
            "Pharmacy_Drug_License": pharmacy["license"],
            "Lead_Pharmacist": pharmacy["leadPharmacist"],

            # Clinical & ICD-10 Coding
            "ICD10_Code": diag_item["icd10"],
            "ICD10_Diagnosis": diag_item["title"],
            "Clinical_Category": diag_item["category"],
            "Prescribed_Medications_Detail": med_summary,
            "Item_Count": len(medications),

            # Cryptographic Security (AES-256-GCM)
            "Encryption_Cipher": "AES-256-GCM",
            "Vault_Transit_Key": f"vault-transit-key-{patient_wallet[:10]}",
            "IV_96bit_Hex": iv_hex,
            "Auth_Tag_128bit_Hex": auth_tag_hex,
            "ZK_Proof_Commitment": "0x" + hashlib.sha256(f"ZK_PROOF_{i}_{patient_short_id}".encode()).hexdigest(),

            # Decentralized Storage & Hedera HCS Ledger
            "IPFS_CID": ipfs_cid,
            "Storage_Provider": "Pinata Dedicated IPFS Gateway",
            "HCS_Topic_ID": "0.0.4891024",
            "HCS_Sequence_Number": hcs_seq,
            "HCS_Consensus_Timestamp": consensus_timestamp,
            "HCS_Running_Hash": running_hash,
            "Hedera_Gas_Fee_HBAR": 0.00010,

            # DPDP Act 2023 Compliance
            "DPDP_Purpose_Limitation": "Pharmacy Dispensation Verification & Medication Fulfillment",
            "Identity_Verification_Status": "Identity Verified (Wallet Protocol)",
            "Blockchain_Audit_Status": "Blockchain Audit Enabled"
        }
        raw_records.append(record)

    df_main = pd.DataFrame(raw_records)

    # Save JSON & CSV
    with open(json_path, 'w', encoding='utf-8') as f:
        json.dump(raw_records, f, indent=2)
    df_main.to_csv(csv_path, index=False, encoding='utf-8')
    df_main.to_csv(root_csv_path, index=False, encoding='utf-8')

    # Summary DataFrames
    status_dist = df_main.groupby('Dispensation_Status').size().reset_index(name='Count')
    status_dist['Percentage'] = (status_dist['Count'] / total_rows * 100).round(2).astype(str) + '%'

    pharm_dist = df_main.groupby('Assigned_Pharmacy').size().reset_index(name='Total_Prescriptions')
    pharm_dist['Total_Revenue_INR'] = df_main.groupby('Assigned_Pharmacy')['Total_Bill_Amount_INR'].sum().round(2).values
    pharm_dist['Avg_Bill_INR'] = df_main.groupby('Assigned_Pharmacy')['Total_Bill_Amount_INR'].mean().round(2).values

    diag_dist = df_main.groupby('Clinical_Category').size().reset_index(name='Prescription_Count')
    diag_dist['Percentage'] = (diag_dist['Prescription_Count'] / total_rows * 100).round(2).astype(str) + '%'

    # Write Excel with Multiple Sheets
    with pd.ExcelWriter(excel_path, engine='openpyxl') as writer:
        df_main.to_excel(writer, sheet_name='Pending_Prescription_Queue', index=False)
        status_dist.to_excel(writer, sheet_name='Fulfillment_Status_Summary', index=False)
        pharm_dist.to_excel(writer, sheet_name='Pharmacy_Revenue_Breakdown', index=False)
        diag_dist.to_excel(writer, sheet_name='Therapeutic_Category_Share', index=False)

    # Write Excel copy to Root
    with pd.ExcelWriter(root_excel_path, engine='openpyxl') as writer:
        df_main.to_excel(writer, sheet_name='Pending_Prescription_Queue', index=False)
        status_dist.to_excel(writer, sheet_name='Fulfillment_Status_Summary', index=False)
        pharm_dist.to_excel(writer, sheet_name='Pharmacy_Revenue_Breakdown', index=False)
        diag_dist.to_excel(writer, sheet_name='Therapeutic_Category_Share', index=False)

    # Style Excel
    for target_file in [excel_path, root_excel_path]:
        wb = openpyxl.load_workbook(target_file)
        header_fill = PatternFill(start_color="0D5C3A", end_color="0D5C3A", fill_type="solid") # Emerald green matching Pharmacy theme
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        for sheet in wb.sheetnames:
            ws = wb[sheet]
            ws.views.sheetView[0].showGridLines = True
            
            # Header
            for col in range(1, ws.max_column + 1):
                c = ws.cell(row=1, column=col)
                c.fill = header_fill
                c.font = header_font
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            # Cells
            for col in ws.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.row > 1:
                        cell.border = border
                        cell.alignment = Alignment(vertical="center")
                    val_str = str(cell.value or '')
                    if len(val_str) > max_len:
                        max_len = len(val_str)
                ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 45)
            ws.row_dimensions[1].height = 28

        wb.save(target_file)

    print(f"[SUCCESS] Generated 800-Row Pharmacy Dataset:")
    print(f" -> Excel (Workspace Root): {root_excel_path}")
    print(f" -> Excel (Scripts Folder): {excel_path}")
    print(f" -> CSV: {csv_path}")
    print(f" -> JSON: {json_path}")
    print("\n--- Summary Breakdown ---")
    print(status_dist.to_string(index=False))

if __name__ == "__main__":
    generate_pharmacy_dataset(800)
