import os
import json
import random
import hashlib
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Seed for reproducibility
random.seed(42)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT_DIR = os.path.abspath(os.path.join(BASE_DIR, '..', '..'))

# Base constants
HCS_TOPIC_ID = "0.0.4891024"
BASE_EPOCH_MS = 1789300000000 # Sept 2026

# Common patient names
FIRST_NAMES = [
    "Aarav", "Aditi", "Ananya", "Arjun", "Deepak", "Divya", "Gautam", "Ishaan", "Kavya", "Manish",
    "Meera", "Neha", "Nikhil", "Pooja", "Pranav", "Priyanka", "Rahul", "Rakesh", "Riya", "Rohan",
    "Sakshi", "Sameer", "Sanjay", "Shreya", "Siddharth", "Sneha", "Tanvi", "Varun", "Vikas", "Yashwanth",
    "Amit", "Anand", "Bhavna", "Chetan", "Geeta", "Harish", "Kiran", "Lata", "Madhav", "Nandini"
]
LAST_NAMES = [
    "Sharma", "Verma", "Patel", "Reddy", "Gupta", "Nair", "Iyer", "Rao", "Deshmukh", "Choudhury",
    "Mehta", "Bhatia", "Joshi", "Kapoor", "Banerjee", "Chatterjee", "Mishra", "Pandey", "Saxena", "Agarwal",
    "Kulkarni", "Shetty", "Pillai", "Menon", "Gowda", "Hegde", "Ghosh", "Mukherjee", "Swaminathan", "Malhotra"
]

# Hospitals
HOSPITALS = [
    {"id": "HOSP-001", "name": "Apollo Multispecialty Hospital", "wallet": "0x155Af6ECaFb48861dA7d16Fb8Af2f6ce9d6DD779", "city": "Hyderabad / Chennai"},
    {"id": "HOSP-002", "name": "Max Super Speciality Cancer Institute", "wallet": "0x28B7eA94B34C8F52D19E50Eb14769Da20F5D54b4", "city": "New Delhi / NCR"},
    {"id": "HOSP-003", "name": "Fortis Memorial Neuro & Trauma Center", "wallet": "0x84f93618C1D541bB1Dcf1373A8eE500A7D439634", "city": "Gurugram / Mumbai"},
    {"id": "HOSP-004", "name": "Care & Cure Pediatric & Maternity Hospital", "wallet": "0x937F51E09477BfD31E56c5e533038a8eEAc096b7", "city": "Bengaluru"},
    {"id": "HOSP-005", "name": "Apex Diagnostic & Pathology Hub", "wallet": "0x04Fee3FD1B338d12FFD6dBD8d66dE1e8e0BB99cB", "city": "Mumbai / Pune"}
]

# Doctors
DOCTORS = [
    {"name": "Dr. Sarah Jenkins", "wallet": "0x04Fee3FD1B338d12FFD6dBD8d66dE1e8e0BB99cB", "hospital": "Apollo Multispecialty Hospital", "regNo": "NMC-48192", "specialty": "Cardiology"},
    {"name": "Dr. Rajesh Sharma", "wallet": "0x3c44CdDdB6a900fa2b585dd299e03d12FA4293BC", "hospital": "Apollo Multispecialty Hospital", "regNo": "NMC-78491", "specialty": "Gastroenterology"},
    {"name": "Dr. Vikram Seth", "wallet": "0x15d34AAf54267DB7D7c367839AAf71A00a2C6A65", "hospital": "Max Super Speciality Cancer Institute", "regNo": "NMC-99014", "specialty": "Oncology"},
    {"name": "Dr. Ananya Roy", "wallet": "0x9965507D1a55bcC2695C58ba16FB37d819B0A4dc", "hospital": "Max Super Speciality Cancer Institute", "regNo": "NMC-51203", "specialty": "Hematology"},
    {"name": "Dr. Rohan Patel", "wallet": "0x976EA74026E726554dB657fA54763abd0C3a0aa9", "hospital": "Fortis Memorial Neuro & Trauma Center", "regNo": "NMC-11029", "specialty": "Neurology"},
    {"name": "Dr. Preeti Deshmukh", "wallet": "0x23618e81E3f5cdF7f54C3d65f7FBc0aBf5B21E8f", "hospital": "Care & Cure Pediatric & Maternity Hospital", "regNo": "NMC-82194", "specialty": "Obstetrics & Pediatrics"},
    {"name": "Dr. Sunita Narang", "wallet": "0xBcd4042DE499D14e55001CcbB24a551F3b954096", "hospital": "Apex Diagnostic & Pathology Hub", "regNo": "NMC-90812", "specialty": "Endocrinology & Internal Medicine"}
]

# Insurance Providers
INSURANCE_COMPANIES = [
    {"id": "INS-001", "name": "Star Health & Allied Insurance", "wallet": "0x70997970C51812dc3A010C7d01b50e0d17dc79C8", "license": "IRDAI-REG-129", "tpa": "MediAssist TPA"},
    {"id": "INS-002", "name": "HDFC ERGO Health Insurance", "wallet": "0x3C44CdDdB6a900fa2b585dd299e03d12FA4293BC", "license": "IRDAI-REG-146", "tpa": "Paramount TPA"},
    {"id": "INS-003", "name": "ICICI Lombard General Insurance", "wallet": "0x90F79bf6EB2c4f870365E785982E1f101E93b906", "license": "IRDAI-REG-115", "tpa": "Vidal Health TPA"},
    {"id": "INS-004", "name": "Care Health Insurance (Religare)", "wallet": "0x15d34AAf54267DB7D7c367839AAf71A00a2C6A65", "license": "IRDAI-REG-148", "tpa": "Heritage Health TPA"},
    {"id": "INS-005", "name": "Niva Bupa Health Insurance", "wallet": "0x9965507D1a55bcC2695C58ba16FB37d819B0A4dc", "license": "IRDAI-REG-145", "tpa": "MDIndia Health TPA"}
]

# Diagnostic Labs
LAB_PROVIDERS = [
    {"id": "LAB-001", "name": "Apex Diagnostic Reference Labs", "wallet": "0x04Fee3FD1B338d12FFD6dBD8d66dE1e8e0BB99cB", "nabl": "NABL-MC-2091", "director": "Dr. Sunita Narang, MD Path"},
    {"id": "LAB-002", "name": "Dr. Lal PathLabs Central Reference Hub", "wallet": "0x976EA74026E726554dB657fA54763abd0C3a0aa9", "nabl": "NABL-MC-1044", "director": "Dr. Arvind Lal, FRCPath"},
    {"id": "LAB-003", "name": "Metropolis Healthcare Molecular Laboratory", "wallet": "0x14dC79964da2C08b23698B3D3cc7Ca32193d9955", "nabl": "NABL-MC-3819", "director": "Dr. Nilesh Shah, MD"},
    {"id": "LAB-004", "name": "Apollo Diagnostics Advanced Enclave", "wallet": "0x23618e81E3f5cdF7f54C3d65f7FBc0aBf5B21E8f", "nabl": "NABL-MC-5102", "director": "Dr. Shanti Swaminathan, MD"},
    {"id": "LAB-005", "name": "SRL Diagnostics Regional Core Lab", "wallet": "0xa0Ee7A142d267C1f36714E4a8F75612F20a79720", "nabl": "NABL-MC-4491", "director": "Dr. Ramesh Sarin, DNB Path"}
]

# Master Clinical Cases
CLINICAL_CASES = [
    {
        "icd10": "I10", "title": "Essential (primary) hypertension", "cat": "Cardiology", "severity": "Moderate",
        "meds": "Amlodipine 5mg QD, Telmisartan 40mg QD",
        "labs": ["Lipid Profile", "Serum Creatinine", "12-Lead ECG"],
        "labResults": {"Lipid Profile": "Total Chol: 215 mg/dL (High), HDL: 42 mg/dL, LDL: 138 mg/dL", "Serum Creatinine": "0.95 mg/dL (Normal)", "12-Lead ECG": "Normal Sinus Rhythm, Voltage criteria for LVH negative"},
        "vitals": {"bp": "148/92 mmHg", "hr": 78, "spo2": 98, "glucose": 110},
        "baseClaimINR": 28500
    },
    {
        "icd10": "I21.9", "title": "Acute myocardial infarction, unspecified", "cat": "Cardiology", "severity": "Emergency",
        "meds": "Ticagrelor 90mg BID, Aspirin 325mg, Atorvastatin 80mg, IV Heparin",
        "labs": ["High-Sensitivity Troponin-I", "CK-MB", "Coronary Angiogram", "2D Echocardiogram"],
        "labResults": {"High-Sensitivity Troponin-I": "2450 ng/L (Critical Alert > 14 ng/L)", "CK-MB": "85 U/L (High)", "2D Echocardiogram": "Anterior wall hypokinesia, LVEF 40%"},
        "vitals": {"bp": "95/60 mmHg", "hr": 112, "spo2": 93, "glucose": 165},
        "baseClaimINR": 245000
    },
    {
        "icd10": "E11.9", "title": "Type 2 diabetes mellitus without complications", "cat": "Endocrinology", "severity": "Moderate",
        "meds": "Metformin 500mg BID, Dapagliflozin 10mg QD",
        "labs": ["HbA1c Glycated Hemoglobin", "Fasting Blood Glucose", "Urine Microalbumin"],
        "labResults": {"HbA1c Glycated Hemoglobin": "7.4% (Elevated > 6.5%)", "Fasting Blood Glucose": "142 mg/dL (High)", "Urine Microalbumin": "18 mg/L (Normal)"},
        "vitals": {"bp": "130/82 mmHg", "hr": 72, "spo2": 99, "glucose": 142},
        "baseClaimINR": 18500
    },
    {
        "icd10": "C50.919", "title": "Malignant neoplasm of female breast", "cat": "Oncology", "severity": "Critical",
        "meds": "Tamoxifen 20mg QD, Trastuzumab 6mg/kg IV Q3W, Paclitaxel",
        "labs": ["ER/PR/HER2 Histopathology", "Mammography Biopsy", "BRCA1/2 Panel", "PET-CT Whole Body"],
        "labResults": {"ER/PR/HER2 Histopathology": "Infiltrating Ductal Carcinoma Grade 2, ER positive (80%), HER2 3+ Positive", "PET-CT Whole Body": "FDG avid primary lesion, no distant metastases"},
        "vitals": {"bp": "122/76 mmHg", "hr": 84, "spo2": 98, "glucose": 105},
        "baseClaimINR": 380000
    },
    {
        "icd10": "G40.909", "title": "Epilepsy, unspecified, without status epilepticus", "cat": "Neurology", "severity": "High",
        "meds": "Levetiracetam 500mg BID, Clobazam 10mg QHS",
        "labs": ["Electroencephalogram (EEG)", "Brain MRI Epilepsy Protocol", "Serum Levetiracetam Level"],
        "labResults": {"Electroencephalogram (EEG)": "Generalized spike and wave discharges, 3Hz paroxysms", "Brain MRI Epilepsy Protocol": "Right mesial temporal sclerosis"},
        "vitals": {"bp": "118/74 mmHg", "hr": 76, "spo2": 99, "glucose": 95},
        "baseClaimINR": 42000
    },
    {
        "icd10": "S06.9X0A", "title": "Unspecified intracranial injury / Concussion", "cat": "Trauma & Emergency", "severity": "Emergency",
        "meds": "Mannitol 20% 100mL IV, Phenytoin 100mg IV, Analgesics",
        "labs": ["Emergency Non-Contrast Brain CT", "C-Spine X-Ray", "Coagulation Profile PT/INR"],
        "labResults": {"Emergency Non-Contrast Brain CT": "Small left frontal contusion, no midline shift, no mass effect", "Coagulation Profile PT/INR": "PT 12.8s, INR 1.05 (Normal)"},
        "vitals": {"bp": "150/95 mmHg", "hr": 64, "spo2": 97, "glucose": 130},
        "baseClaimINR": 115000
    },
    {
        "icd10": "K21.9", "title": "Gastro-esophageal reflux disease", "cat": "Gastroenterology", "severity": "Low",
        "meds": "Pantoprazole 40mg QD pre-breakfast, Domperidone 30mg QD",
        "labs": ["Upper Gastrointestinal Endoscopy", "H. Pylori Stool Antigen Test"],
        "labResults": {"Upper Gastrointestinal Endoscopy": "Los Angeles Grade B Reflux Esophagitis", "H. Pylori Stool Antigen Test": "Negative"},
        "vitals": {"bp": "120/80 mmHg", "hr": 70, "spo2": 99, "glucose": 98},
        "baseClaimINR": 12500
    },
    {
        "icd10": "A90", "title": "Dengue fever [classical dengue]", "cat": "Infectious Disease", "severity": "High",
        "meds": "Oral Rehydration Salts 1L/day, Paracetamol 650mg SOS",
        "labs": ["Dengue NS1 Antigen ELISA", "Complete Blood Count Platelet Serial", "Liver Function Tests"],
        "labResults": {"Dengue NS1 Antigen ELISA": "Positive", "Complete Blood Count Platelet Serial": "Platelets: 45,000 /uL (Critical Low), Hematocrit: 44%", "Liver Function Tests": "SGOT: 120 U/L, SGPT: 98 U/L (Elevated)"},
        "vitals": {"bp": "100/68 mmHg", "hr": 96, "spo2": 98, "glucose": 102},
        "baseClaimINR": 32000
    }
]

# Helper styling function
def style_workbook(wb, header_color_hex="1F4E79"):
    header_fill = PatternFill(start_color=header_color_hex, end_color=header_color_hex, fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    for sheet in wb.sheetnames:
        ws = wb[sheet]
        ws.views.sheetView[0].showGridLines = True
        
        for col_idx in range(1, ws.max_column + 1):
            c = ws.cell(row=1, column=col_idx)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.row > 1:
                    cell.border = border
                    cell.alignment = Alignment(vertical="center")
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = min(max(max_len + 4, 14), 45)
        ws.row_dimensions[1].height = 28


# ==============================================================================
# 1. GENERATE INSURANCE DATASET (1,000 ROWS)
# ==============================================================================
def generate_insurance_dataset(n=1000):
    print("Generating Insurance Claims & Policy Adjudication Dataset (1,000 rows)...")
    records = []

    for i in range(1, n + 1):
        patient_name = f"{random.choice(FIRST_NAMES)} {random.choice(LAST_NAMES)}"
        patient_short_id = str(849200 + (i % 600))
        patient_wallet = "0x" + hashlib.sha256(f"PATIENT_INS_{patient_short_id}".encode()).hexdigest()[:40]

        ins = random.choice(INSURANCE_COMPANIES)
        hospital = random.choice(HOSPITALS)
        doctor = random.choice(DOCTORS)
        case = random.choice(CLINICAL_CASES)

        policy_no = f"POL-{ins['name'][:4].upper()}-{patient_short_id}-{random.randint(100, 999)}"
        claim_id = f"CLM-{str(i).zfill(5)}"
        
        claim_amount = round(case["baseClaimINR"] * random.uniform(0.9, 1.35), 2)
        deductible = round(min(5000.0, claim_amount * 0.05), 2)
        co_pay_pct = random.choice([0, 10, 15, 20])

        # Claim status distribution matching Portal
        status_rand = random.random()
        if status_rand < 0.60:
            claim_status = "Approved & Settled"
            approved_amount = round(claim_amount - deductible - (claim_amount * (co_pay_pct / 100)), 2)
            onchain_status = "Ledger Verified & Payout Secured"
            consent_status = "Active - Full Disclosure"
        elif status_rand < 0.82:
            claim_status = "Pending Adjudication"
            approved_amount = 0.0
            onchain_status = "Audit Logged - Under Verification"
            consent_status = "Active - Time-Bound Access"
        elif status_rand < 0.93:
            claim_status = "Awaiting Patient Consent Approval"
            approved_amount = 0.0
            onchain_status = "Access Requested via Ledger"
            consent_status = "Pending Patient Authorization"
        else:
            claim_status = "Rejected - Policy Exclusion"
            approved_amount = 0.0
            onchain_status = "Rejected with Ledger Audit Reason"
            consent_status = "Expired / Revoked"

        requested_data_types = random.choice([
            "Lab Reports; Diagnosis; Treatment History; Discharge Summary",
            "Diagnosis; Lab Reports; Itemized Pharmacy Bill",
            "Discharge Summary; Surgical Notes; Lab Reports",
            "Emergency Admission Record; Diagnostic Imaging CT/MRI"
        ])

        hcs_seq = 3000 + i
        time_offset = (i * 1920) + random.randint(100, 800)
        consensus_timestamp = pd.to_datetime(BASE_EPOCH_MS + time_offset, unit='ms').strftime('%Y-%m-%d %H:%M:%S UTC')
        ipfs_cid = "Qm" + hashlib.sha256(f"CLAIM_IPFS_{i}_{claim_id}".encode()).hexdigest()[:44]
        running_hash = "0x" + hashlib.sha256(f"HCS_INS_{hcs_seq}_{ipfs_cid}".encode()).hexdigest()

        records.append({
            "Claim_ID": claim_id,
            "Policy_Number": policy_no,
            "Insurance_Provider": ins["name"],
            "Insurance_Wallet": ins["wallet"],
            "IRDAI_License": ins["license"],
            "Third_Party_Administrator": ins["tpa"],
            "Patient_Name": patient_name,
            "Patient_Short_ID": patient_short_id,
            "Patient_Wallet": patient_wallet,
            "Treating_Hospital": hospital["name"],
            "Hospital_Wallet": hospital["wallet"],
            "Attending_Physician": doctor["name"],
            "Physician_Registration": doctor["regNo"],
            "ICD10_Code": case["icd10"],
            "Diagnosis_Title": case["title"],
            "Clinical_Category": case["cat"],
            "Claim_Amount_INR": claim_amount,
            "Approved_Amount_INR": approved_amount,
            "Deductible_INR": deductible,
            "Co_Pay_Percentage": f"{co_pay_pct}%",
            "Claim_Status": claim_status,
            "OnChain_Verification_Status": onchain_status,
            "DPDP_Consent_Status": consent_status,
            "Requested_Data_Types": requested_data_types,
            "Consent_Purpose_Limitation": "Insurance Claim Adjudication & Pre-Auth Verification",
            "Consent_Duration_Window": "24 Hours (Statutory Emergency Window)",
            "Encryption_Cipher": "AES-256-GCM",
            "Vault_Transit_Key_Handle": f"vault-transit-key-{patient_wallet[:10]}",
            "IPFS_Claim_Document_CID": ipfs_cid,
            "HCS_Topic_ID": HCS_TOPIC_ID,
            "HCS_Sequence_Number": hcs_seq,
            "HCS_Consensus_Timestamp": consensus_timestamp,
            "HCS_Running_Hash": running_hash,
            "Hedera_Gas_Fee_HBAR": 0.00010
        })

    df = pd.DataFrame(records)
    
    # Save CSV & JSON
    df.to_csv(os.path.join(BASE_DIR, 'insurance_dataset_1000.csv'), index=False, encoding='utf-8')
    df.to_csv(os.path.join(ROOT_DIR, 'insurance_dataset_1000.csv'), index=False, encoding='utf-8')
    with open(os.path.join(BASE_DIR, 'insurance_dataset_1000.json'), 'w', encoding='utf-8') as f:
        json.dump(records, f, indent=2)

    # Summaries
    status_summary = df.groupby('Claim_Status').size().reset_index(name='Claim_Count')
    status_summary['Total_Claim_INR'] = df.groupby('Claim_Status')['Claim_Amount_INR'].sum().round(2).values
    status_summary['Total_Approved_INR'] = df.groupby('Claim_Status')['Approved_Amount_INR'].sum().round(2).values

    ins_summary = df.groupby('Insurance_Provider').size().reset_index(name='Total_Claims')
    ins_summary['Total_Claim_Volume_INR'] = df.groupby('Insurance_Provider')['Claim_Amount_INR'].sum().round(2).values

    # Excel
    for out_path in [os.path.join(BASE_DIR, 'insurance_dataset_1000.xlsx'), os.path.join(ROOT_DIR, 'insurance_dataset_1000.xlsx')]:
        with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Insurance_Claims_Ledger', index=False)
            status_summary.to_excel(writer, sheet_name='Claim_Adjudication_Summary', index=False)
            ins_summary.to_excel(writer, sheet_name='Payer_Volume_Breakdown', index=False)
        wb = openpyxl.load_workbook(out_path)
        style_workbook(wb, "1B365D") # Navy blue
        wb.save(out_path)

    print(f" -> Insurance Dataset Saved: {os.path.join(ROOT_DIR, 'insurance_dataset_1000.xlsx')}")


# ==============================================================================
# 2. GENERATE DOCTOR DATASET (1,000 ROWS)
# ==============================================================================
def generate_doctor_dataset(n=1000):
    print("Generating Doctor Clinical Encounters & Prescriptions Dataset (1,000 rows)...")
    records = []

    for i in range(1, n + 1):
        patient_name = f"{random.choice(FIRST_NAMES)} {random.choice(LAST_NAMES)}"
        patient_short_id = str(849200 + (i % 600))
        patient_wallet = "0x" + hashlib.sha256(f"PATIENT_DOC_{patient_short_id}".encode()).hexdigest()[:40]
        patient_age = 18 + (i % 72)
        patient_gender = "female" if (i % 2 == 0) else "male"
        blood_group = random.choice(["O+", "A+", "B+", "AB+", "O-", "A-", "B-", "AB-"])

        doctor = random.choice(DOCTORS)
        hospital = random.choice(HOSPITALS)
        case = random.choice(CLINICAL_CASES)

        encounter_id = f"ENC-{str(i).zfill(5)}"
        
        # Break glass logic matching Doctor portal
        is_emergency = case["severity"] == "Emergency"
        break_glass_invoked = "YES" if (is_emergency and random.random() < 0.70) else "NO"
        
        if break_glass_invoked == "YES":
            consent_status = "Break-Glass Emergency Invoked (Active Override)"
            break_glass_reason = f"Acute Life-Threatening Presentation: {case['title']}"
            ephemeral_ttl = "60 Minutes (Auto-Expiring Ephemeral Key)"
            audit_note = "Emergency access granted; immediate HCS compliance event broadcasted"
        else:
            break_glass_reason = "N/A"
            ephemeral_ttl = "N/A"
            if random.random() < 0.75:
                consent_status = "Active On-Chain Consent (Granted by Patient)"
                audit_note = "Direct patient consent validated via ConsentManager.sol"
            else:
                consent_status = "Pending Patient Authorization"
                audit_note = "Access request pending approval from patient wallet"

        # AI Clinical Safety Audit matching DoctorDashboard.jsx
        ai_safety_options = [
            "Passed - Zero Drug-Drug Interaction Warning (Confidence 99.4%)",
            "Advisory: Monitor Serum Potassium & Renal Function with ACEi/ARB",
            "Advisory: Confirm Patient Creatinine Clearance prior to Metformin XR",
            "Passed - Safe pediatric/geriatric dosage tier verified"
        ]
        ai_safety = random.choice(ai_safety_options)

        hcs_seq = 4000 + i
        time_offset = (i * 2100) + random.randint(100, 900)
        consensus_timestamp = pd.to_datetime(BASE_EPOCH_MS + time_offset, unit='ms').strftime('%Y-%m-%d %H:%M:%S UTC')
        ipfs_cid = "Qm" + hashlib.sha256(f"DOC_ENCOUNTER_IPFS_{i}_{encounter_id}".encode()).hexdigest()[:44]
        running_hash = "0x" + hashlib.sha256(f"HCS_DOC_{hcs_seq}_{ipfs_cid}".encode()).hexdigest()

        records.append({
            "Encounter_ID": encounter_id,
            "Patient_Name": patient_name,
            "Patient_Short_ID": patient_short_id,
            "Patient_Wallet": patient_wallet,
            "Patient_Age": patient_age,
            "Patient_Gender": patient_gender,
            "Blood_Group": blood_group,
            "Attending_Physician": doctor["name"],
            "Physician_Wallet": doctor["wallet"],
            "NMC_Registration_No": doctor["regNo"],
            "Medical_Specialty": doctor["specialty"],
            "Hospital_Affiliation": hospital["name"],
            "Hospital_Wallet": hospital["wallet"],
            "Chief_Complaint_Diagnosis": case["title"],
            "ICD10_Diagnosis_Code": case["icd10"],
            "Clinical_Category": case["cat"],
            "Severity_Tier": case["severity"],
            "Blood_Pressure_mmHg": case["vitals"]["bp"],
            "Heart_Rate_BPM": case["vitals"]["hr"],
            "SpO2_Percentage": f"{case['vitals']['spo2']}%",
            "Blood_Glucose_mg_dL": case["vitals"]["glucose"],
            "Prescribed_Medication_Regimen": case["meds"],
            "Prescription_Sensitivity": random.choice(["Low", "Medium", "High"]),
            "AI_Clinical_Safety_Check": ai_safety,
            "FHIR_Resource_Bundle": f"FHIR_R4_Bundle_{case['icd10']}_Condition_MedicationRequest",
            "Consent_Access_Scope": "All Clinical Records & Diagnostic Evidence",
            "DPDP_Consent_Status": consent_status,
            "Break_Glass_Emergency_Invoked": break_glass_invoked,
            "Break_Glass_Clinical_Justification": break_glass_reason,
            "Ephemeral_Key_TTL": ephemeral_ttl,
            "Blockchain_Audit_Note": audit_note,
            "Encryption_Cipher": "AES-256-GCM",
            "Vault_Transit_Key_Handle": f"vault-transit-key-{patient_wallet[:10]}",
            "IPFS_Encounter_CID": ipfs_cid,
            "HCS_Topic_ID": HCS_TOPIC_ID,
            "HCS_Sequence_Number": hcs_seq,
            "HCS_Consensus_Timestamp": consensus_timestamp,
            "HCS_Running_Hash": running_hash,
            "Hedera_Gas_Fee_HBAR": 0.00010
        })

    df = pd.DataFrame(records)

    # Save CSV & JSON
    df.to_csv(os.path.join(BASE_DIR, 'doctor_dataset_1000.csv'), index=False, encoding='utf-8')
    df.to_csv(os.path.join(ROOT_DIR, 'doctor_dataset_1000.csv'), index=False, encoding='utf-8')
    with open(os.path.join(BASE_DIR, 'doctor_dataset_1000.json'), 'w', encoding='utf-8') as f:
        json.dump(records, f, indent=2)

    # Summaries
    specialty_summary = df.groupby('Medical_Specialty').size().reset_index(name='Total_Consultations')
    break_glass_summary = df.groupby('Break_Glass_Emergency_Invoked').size().reset_index(name='Encounter_Count')
    severity_summary = df.groupby('Severity_Tier').size().reset_index(name='Patient_Count')

    # Excel
    for out_path in [os.path.join(BASE_DIR, 'doctor_dataset_1000.xlsx'), os.path.join(ROOT_DIR, 'doctor_dataset_1000.xlsx')]:
        with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Doctor_Clinical_Encounters', index=False)
            specialty_summary.to_excel(writer, sheet_name='Specialty_Workload', index=False)
            break_glass_summary.to_excel(writer, sheet_name='Break_Glass_Distribution', index=False)
            severity_summary.to_excel(writer, sheet_name='Acuity_Severity_Tiers', index=False)
        wb = openpyxl.load_workbook(out_path)
        style_workbook(wb, "4A2E7A") # Royal purple
        wb.save(out_path)

    print(f" -> Doctor Dataset Saved: {os.path.join(ROOT_DIR, 'doctor_dataset_1000.xlsx')}")


# ==============================================================================
# 3. GENERATE LAB DATASET (1,000 ROWS)
# ==============================================================================
def generate_lab_dataset(n=1000):
    print("Generating Diagnostic Pathology & Radiology Laboratory Dataset (1,000 rows)...")
    records = []

    LAB_TEST_CATALOG = [
        {"cat": "Blood Tests", "test": "Complete Blood Count (CBC)", "specimen": "Whole Blood (EDTA)", "cost": 450, "refRange": "Hb: 12-16 g/dL, WBC: 4000-11000 /uL, Platelets: 1.5-4.5 Lakhs/uL", "sampleResult": "Hb: 13.5 g/dL, WBC: 7200 /uL, Platelets: 2.4 Lakhs/uL", "flag": "NORMAL"},
        {"cat": "Blood Tests", "test": "Thrombocytopenia Platelet Monitoring", "specimen": "Whole Blood (EDTA)", "cost": 350, "refRange": "Platelets: 150,000 - 450,000 /uL", "sampleResult": "Platelets: 42,000 /uL (Severe Low)", "flag": "CRITICAL_ALERT"},
        {"cat": "Metabolic Panel", "test": "Comprehensive Metabolic Panel (CMP)", "specimen": "Serum", "cost": 1200, "refRange": "Na: 135-145, K: 3.5-5.0, Creatinine: 0.7-1.3 mg/dL", "sampleResult": "Na: 138 mmol/L, K: 4.1 mmol/L, Creat: 1.05 mg/dL", "flag": "NORMAL"},
        {"cat": "Lipid Profile", "test": "Comprehensive Lipid Panel", "specimen": "Serum (Fasting 12h)", "cost": 850, "refRange": "Chol: <200 mg/dL, HDL: >40 mg/dL, LDL: <100 mg/dL, TG: <150 mg/dL", "sampleResult": "Total Chol: 235 mg/dL (High), LDL: 155 mg/dL, TG: 180 mg/dL", "flag": "HIGH"},
        {"cat": "Liver Function Tests", "test": "Liver Function Profile (LFT)", "specimen": "Serum", "cost": 950, "refRange": "ALT: 7-56 U/L, AST: 10-40 U/L, Bilirubin: 0.1-1.2 mg/dL", "sampleResult": "ALT: 88 U/L (High), AST: 74 U/L (High), Total Bili: 1.8 mg/dL", "flag": "HIGH"},
        {"cat": "Kidney Function Tests", "test": "Renal Function Profile (KFT)", "specimen": "Serum", "cost": 850, "refRange": "BUN: 7-20 mg/dL, Creatinine: 0.6-1.2 mg/dL, eGFR: >60 mL/min", "sampleResult": "BUN: 14 mg/dL, Creatinine: 0.9 mg/dL, eGFR: 92 mL/min", "flag": "NORMAL"},
        {"cat": "Endocrinology Tests", "test": "HbA1c Glycated Hemoglobin", "specimen": "Whole Blood (EDTA)", "cost": 650, "refRange": "<5.7% Normal, 5.7-6.4% Prediabetes, >=6.5% Diabetes", "sampleResult": "HbA1c: 7.6% (Diabetic Range)", "flag": "HIGH"},
        {"cat": "Endocrinology Tests", "test": "Thyroid Function Panel (TSH, Free T3, Free T4)", "specimen": "Serum", "cost": 750, "refRange": "TSH: 0.45-4.5 uIU/mL, FT3: 2.0-4.4 pg/mL, FT4: 0.8-1.8 ng/dL", "sampleResult": "TSH: 6.8 uIU/mL (Hypothyroid), FT4: 0.72 ng/dL (Low)", "flag": "LOW"},
        {"cat": "Cardiology Tests", "test": "High-Sensitivity Cardiac Troponin-I (hs-cTnI)", "specimen": "Serum / Plasma", "cost": 1800, "refRange": "< 14 ng/L (99th Percentile Upper Reference Limit)", "sampleResult": "hs-cTnI: 1850 ng/L (Markedly Elevated Acute Myocardial Infarction)", "flag": "CRITICAL_ALERT"},
        {"cat": "Microbiology Tests", "test": "Dengue Duo NS1 Ag + IgM/IgG Antibody ELISA", "specimen": "Serum", "cost": 1400, "refRange": "Non-Reactive / Negative", "sampleResult": "Dengue NS1 Antigen: POSITIVE, IgM: Reactive", "flag": "POSITIVE_ABNORMAL"},
        {"cat": "Pathology / Biopsy", "test": "Core Needle Biopsy Histopathology & Immunohistochemistry", "specimen": "Tissue Biopsy in 10% Formalin", "cost": 4500, "refRange": "Benign / No malignancy", "sampleResult": "Infiltrating Ductal Carcinoma Grade 2, ER 85%+, PR 60%+, HER2/neu Negative", "flag": "MALIGNANT_POSITIVE"},
        {"cat": "Imaging / Radiology", "test": "High-Resolution Non-Contrast Brain CT Scan", "specimen": "Radiological DICOM Imaging", "cost": 3200, "refRange": "No acute intracranial hemorrhage or midline shift", "sampleResult": "Small left temporal hemorrhagic contusion, no mass effect", "flag": "ABNORMAL_FINDING"},
        {"cat": "Imaging / Radiology", "test": "Whole Body PET-CT Oncology Scan", "specimen": "18F-FDG Radiotracer DICOM", "cost": 18500, "refRange": "Physiological FDG biodistribution, no hypermetabolic malignancy", "sampleResult": "Hypermetabolic primary lesion, SUVmax 9.2, no distant metastases", "flag": "ABNORMAL_FINDING"}
    ]

    for i in range(1, n + 1):
        patient_name = f"{random.choice(FIRST_NAMES)} {random.choice(LAST_NAMES)}"
        patient_short_id = str(849200 + (i % 600))
        patient_wallet = "0x" + hashlib.sha256(f"PATIENT_LAB_{patient_short_id}".encode()).hexdigest()[:40]
        patient_age = 16 + (i % 75)
        patient_gender = "female" if (i % 2 == 0) else "male"

        lab = random.choice(LAB_PROVIDERS)
        doctor = random.choice(DOCTORS)
        hospital = random.choice(HOSPITALS)
        test_item = random.choice(LAB_TEST_CATALOG)

        report_id = f"LAB-REP-{str(i).zfill(5)}"
        accession_no = f"ACC-{lab['id']}-{str(i).zfill(6)}"

        # Result variations
        flag = test_item["flag"]
        result_text = test_item["sampleResult"]
        if flag == "NORMAL" and random.random() < 0.15:
            flag = "BORDERLINE"
            result_text = result_text + " (Borderline clinical correlation advised)"

        sensitivity = "High - Genomic/Oncology" if ("Biopsy" in test_item["test"] or "PET-CT" in test_item["test"]) else ("Medium" if test_item["cost"] > 1000 else "Low")

        # Consent status
        consent_status = "Consent Verified (Active on Blockchain)" if (random.random() < 0.85) else "Emergency Diagnostic Upload (Pending Verification)"

        hcs_seq = 5000 + i
        time_offset = (i * 2250) + random.randint(100, 950)
        consensus_timestamp = pd.to_datetime(BASE_EPOCH_MS + time_offset, unit='ms').strftime('%Y-%m-%d %H:%M:%S UTC')
        ipfs_cid = "Qm" + hashlib.sha256(f"LAB_REPORT_IPFS_{i}_{report_id}".encode()).hexdigest()[:44]
        running_hash = "0x" + hashlib.sha256(f"HCS_LAB_{hcs_seq}_{ipfs_cid}".encode()).hexdigest()

        records.append({
            "Report_ID": report_id,
            "Accession_Number": accession_no,
            "Diagnostic_Laboratory": lab["name"],
            "Laboratory_Wallet": lab["wallet"],
            "NABL_Accreditation_No": lab["nabl"],
            "Laboratory_Director": lab["director"],
            "Patient_Name": patient_name,
            "Patient_Short_ID": patient_short_id,
            "Patient_Wallet": patient_wallet,
            "Patient_Age": patient_age,
            "Patient_Gender": patient_gender,
            "Referring_Physician": doctor["name"],
            "Referring_Doctor_Wallet": doctor["wallet"],
            "Referring_Hospital": hospital["name"],
            "Test_Category": test_item["cat"],
            "Diagnostic_Test_Name": test_item["test"],
            "Specimen_Type": test_item["specimen"],
            "Quantitative_Result_Value": result_text,
            "Biological_Reference_Interval": test_item["refRange"],
            "Diagnostic_Clinical_Flag": flag,
            "Report_Sensitivity_Tier": sensitivity,
            "Laboratory_Bill_Amount_INR": test_item["cost"],
            "DPDP_Consent_Scope": "Diagnostic Laboratory Reports",
            "Consent_Verification_Status": consent_status,
            "Encryption_Cipher": "AES-256-GCM",
            "Vault_Transit_Key_Handle": f"vault-transit-key-{patient_wallet[:10]}",
            "IPFS_Report_CID": ipfs_cid,
            "Storage_Provider": "Pinata Dedicated IPFS Gateway",
            "HCS_Topic_ID": HCS_TOPIC_ID,
            "HCS_Sequence_Number": hcs_seq,
            "HCS_Consensus_Timestamp": consensus_timestamp,
            "HCS_Running_Hash": running_hash,
            "Hedera_Gas_Fee_HBAR": 0.00010
        })

    df = pd.DataFrame(records)

    # Save CSV & JSON
    df.to_csv(os.path.join(BASE_DIR, 'lab_dataset_1000.csv'), index=False, encoding='utf-8')
    df.to_csv(os.path.join(ROOT_DIR, 'lab_dataset_1000.csv'), index=False, encoding='utf-8')
    with open(os.path.join(BASE_DIR, 'lab_dataset_1000.json'), 'w', encoding='utf-8') as f:
        json.dump(records, f, indent=2)

    # Summaries
    category_summary = df.groupby('Test_Category').size().reset_index(name='Total_Tests')
    category_summary['Total_Revenue_INR'] = df.groupby('Test_Category')['Laboratory_Bill_Amount_INR'].sum().values

    flag_summary = df.groupby('Diagnostic_Clinical_Flag').size().reset_index(name='Count')
    flag_summary['Percentage'] = (flag_summary['Count'] / n * 100).round(2).astype(str) + '%'

    lab_summary = df.groupby('Diagnostic_Laboratory').size().reset_index(name='Tests_Conducted')
    lab_summary['Total_Billing_INR'] = df.groupby('Diagnostic_Laboratory')['Laboratory_Bill_Amount_INR'].sum().values

    # Excel
    for out_path in [os.path.join(BASE_DIR, 'lab_dataset_1000.xlsx'), os.path.join(ROOT_DIR, 'lab_dataset_1000.xlsx')]:
        with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Diagnostic_Lab_Reports', index=False)
            category_summary.to_excel(writer, sheet_name='Test_Category_Revenue', index=False)
            flag_summary.to_excel(writer, sheet_name='Clinical_Flag_Distribution', index=False)
            lab_summary.to_excel(writer, sheet_name='Laboratory_Volume_Breakdown', index=False)
        wb = openpyxl.load_workbook(out_path)
        style_workbook(wb, "006666") # Teal/Cyan medical lab
        wb.save(out_path)

    print(f" -> Lab Dataset Saved: {os.path.join(ROOT_DIR, 'lab_dataset_1000.xlsx')}")


if __name__ == "__main__":
    print("================================================================================")
    print("OJASRAKSHA HEALTHCARE ECOSYSTEM DATASET GENERATOR (3,000 ROWS ACROSS 3 PORTALS)")
    print("================================================================================\n")
    generate_insurance_dataset(1000)
    generate_doctor_dataset(1000)
    generate_lab_dataset(1000)
    print("\n[ALL DATASETS GENERATED SUCCESSFULLY]")
